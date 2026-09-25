"""Bounded CSV / XLSX parsing for uploaded data files (sandbox worker only).

* rows, columns and cell sizes are capped; parsing streams rows instead of
  loading whole sheets into memory;
* XLSX is opened read-only with ``data_only=True`` - formulas are *not*
  evaluated (their cached values are read), external links are not loaded and
  macros are never kept; ``defusedxml`` is installed, which openpyxl uses to
  block XML entity attacks (billion laughs / XXE).
"""

from __future__ import annotations

import csv
import math
import zipfile
import zlib
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from nexusflow.core.errors import InvalidInputError
from nexusflow.domain.sources.model import FileUploadConfig

MAX_CELL_CHARS = 10_000
# Blank rows are skipped, but read-only openpyxl produces one for every row
# number a sheet skips: a 5 KB file with a row numbered 1,048,576 would walk a
# million rows. At most this many blank rows are walked beyond the row cap;
# stopping there marks the file truncated, like any other cut.
BLANK_ROW_ALLOWANCE = 1_000
# What openpyxl and the XML parser raise for a workbook that is broken or
# hostile (defusedxml's entity refusals are ValueErrors, XML syntax errors are
# SyntaxErrors, a missing part is a KeyError or IndexError).
_UNREADABLE = (
    ValueError,
    SyntaxError,
    KeyError,
    IndexError,
    TypeError,
    OverflowError,
    EOFError,
    zipfile.BadZipFile,
    zlib.error,
    InvalidFileException,  # not a workbook openpyxl can open
)


@dataclass(frozen=True, slots=True)
class ParsedFile:
    items: list[dict[str, Any]]
    truncated: bool
    rows: int


def _require_columns(header: list[str], config: FileUploadConfig) -> dict[str, int]:
    positions = {name.strip(): index for index, name in enumerate(header)}
    missing = [column for column in config.column_mapping.values() if column not in positions]
    if missing:
        raise InvalidInputError(
            f"Missing columns: {', '.join(sorted(missing))[:300]}", code="upload_missing_columns"
        )
    return {field: positions[column] for field, column in config.column_mapping.items()}


def _cell(value: Any) -> Any:
    if isinstance(value, float) and not math.isfinite(value):
        # 1E999 reads as infinity, which no JSON consumer accepts.
        raise InvalidInputError(
            "A cell holds a number that is not finite.", code="upload_malformed"
        )
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        return value[:MAX_CELL_CHARS]
    return value


def parse_csv(
    path: Path, config: FileUploadConfig, *, max_rows: int, max_columns: int
) -> ParsedFile:
    items: list[dict[str, Any]] = []
    rows = 0
    truncated = False
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=config.delimiter, strict=True)
        try:
            header = next(reader)
        except StopIteration as exc:
            raise InvalidInputError("The file is empty.", code="upload_empty") from exc
        except csv.Error as exc:
            raise InvalidInputError("Malformed CSV header.", code="upload_malformed") from exc
        if len(header) > max_columns:
            raise InvalidInputError("Too many columns.", code="upload_too_many_columns")
        positions = _require_columns(header, config)
        try:
            for row in reader:
                if rows >= max_rows:
                    truncated = True
                    break
                if len(row) > max_columns:
                    raise InvalidInputError(
                        f"Too many columns in row {rows + 2}.", code="upload_too_many_columns"
                    )
                rows += 1
                items.append(
                    {f: _cell(row[i]) if i < len(row) else None for f, i in positions.items()}
                )
        except csv.Error as exc:
            raise InvalidInputError(
                f"Malformed CSV near row {rows + 1}.", code="upload_malformed"
            ) from exc
    return ParsedFile(items=items, truncated=truncated, rows=rows)


def parse_xlsx(
    path: Path, config: FileUploadConfig, *, max_rows: int, max_columns: int
) -> ParsedFile:
    """Parse a workbook; a broken or hostile one is an upload error, not a crash.

    A crash would be retried - parsing the hostile file again - and reported as
    a generic sandbox error.
    """
    try:
        return _parse_xlsx(path, config, max_rows=max_rows, max_columns=max_columns)
    except InvalidInputError:
        raise
    except _UNREADABLE as exc:
        raise InvalidInputError("The workbook could not be read.", code="upload_malformed") from exc


def _parse_xlsx(
    path: Path, config: FileUploadConfig, *, max_rows: int, max_columns: int
) -> ParsedFile:
    # An open file, not the path: openpyxl refuses a path without an Excel
    # extension, and the sandbox stores its input as plain "input".
    with path.open("rb") as handle:
        workbook = openpyxl.load_workbook(
            handle, read_only=True, data_only=True, keep_links=False, keep_vba=False
        )
        try:
            sheet: Any
            if config.sheet_name is not None:
                if config.sheet_name not in workbook.sheetnames:
                    raise InvalidInputError(
                        "The worksheet was not found.", code="upload_missing_sheet"
                    )
                sheet = workbook[config.sheet_name]
            else:
                sheet = workbook.worksheets[0]
            # The header, the capped data rows and the blank-row allowance: openpyxl
            # stops at this row number, so no row numbering makes it walk further.
            last_row = 1 + max_rows + BLANK_ROW_ALLOWANCE
            row_iter = sheet.iter_rows(values_only=True, max_col=max_columns, max_row=last_row)
            header_row = next(row_iter, None)
            if header_row is None:
                raise InvalidInputError("The worksheet is empty.", code="upload_empty")
            header = [str(value).strip() if value is not None else "" for value in header_row]
            positions = _require_columns(header, config)
            items: list[dict[str, Any]] = []
            rows = 0
            walked = 1
            truncated = False
            for row in row_iter:
                walked += 1
                if rows >= max_rows:
                    truncated = True
                    break
                if all(value is None for value in row):
                    continue
                rows += 1
                items.append(
                    {f: _cell(row[i]) if i < len(row) else None for f, i in positions.items()}
                )
            if walked >= last_row:
                truncated = True  # stopped at the walking bound: rows beyond were not read
            return ParsedFile(items=items, truncated=truncated, rows=rows)
        finally:
            workbook.close()


def parse_upload(
    path: Path, config: FileUploadConfig, *, max_rows: int, max_columns: int
) -> ParsedFile:
    if config.format == "csv":
        return parse_csv(path, config, max_rows=max_rows, max_columns=max_columns)
    return parse_xlsx(path, config, max_rows=max_rows, max_columns=max_columns)
