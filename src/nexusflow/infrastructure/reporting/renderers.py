"""Report renderers: JSON, CSV, XLSX and PDF.

Output encoding is context specific:
* CSV cells beginning with formula characters are prefixed with ``'`` (OWASP);
* XLSX cells are written as *string* cells - never formulas - regardless of content;
* PDF text is XML-escaped before it reaches ReportLab's paragraph markup parser,
  so scraped content cannot inject markup (fonts, links, images).
"""

from __future__ import annotations

import csv
import io
import json
import math
from dataclasses import asdict
from typing import Any, cast
from xml.sax.saxutils import escape  # nosec B406

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from nexusflow.core.text import clean_text, spreadsheet_safe, truncate
from nexusflow.domain.catalog.service import describe_field_change
from nexusflow.domain.reports.model import ReportContent, ReportFormat

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_CHANGE_COLUMNS = ("detected_at", "record_key", "change_type", "significance", "score", "changes")
_UNUSUAL_COLUMNS = ["Day", "Changes", "Typical day (median)", "Robust z-score"]
_BAR = colors.HexColor("#2563eb")
_BAR_UNUSUAL = colors.HexColor("#dc2626")


def _diff_summary(diff: Any) -> str:
    if not isinstance(diff, dict):
        return ""
    described = (describe_field_change(name, entry) for name, entry in list(diff.items())[:6])
    return "; ".join(part for part in described if part is not None)


def _row(change: dict[str, Any]) -> list[Any]:
    return [
        change.get("detected_at", ""),
        change.get("record_key", ""),
        change.get("change_type", ""),
        change.get("significance", ""),
        change.get("score", 0),
        _diff_summary(change.get("diff")),
    ]


def render_json(content: ReportContent) -> bytes:
    return json.dumps(asdict(content), default=str, ensure_ascii=False, indent=2).encode("utf-8")


def _unusual_rows(content: ReportContent) -> list[list[Any]]:
    return [
        [a["date"], a["changes"], a["baseline"], a.get("score")] for a in content.volume_anomalies
    ]


def render_csv(content: ReportContent) -> bytes:
    """The detected changes, one per row: CSV carries one table (JSON/XLSX/PDF carry all)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(_CHANGE_COLUMNS)
    for change in content.changes:
        writer.writerow(
            [spreadsheet_safe(str(v)) if isinstance(v, str) else v for v in _row(change)]
        )
    return buffer.getvalue().encode("utf-8-sig")


def _write_rows(sheet: Worksheet, header: list[str], rows: list[list[Any]]) -> None:
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([truncate(str(v), 32_000) if isinstance(v, str) else v for v in row])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_START):
                cell.data_type = "s"  # force a literal string cell, never a formula


def render_xlsx(content: ReportContent) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    if summary is None:  # pragma: no cover - a new workbook always has a sheet
        summary = workbook.create_sheet()
    summary.title = "Summary"
    rows: list[list[Any]] = [
        ["Report", content.title],
        ["Organization", content.organization],
        ["Project", content.project],
        ["Dataset", content.dataset or "all"],
        ["Period", f"{content.period_start:%Y-%m-%d} - {content.period_end:%Y-%m-%d}"],
        ["Generated", content.generated_at.isoformat()],
        ["Executive summary", content.executive_summary],
        ["Trend", content.trend_note],
        ["Unusual days", len(content.volume_anomalies)],
        *[[f"Total {k}", v] for k, v in content.totals.items()],
        *[[f"Significance {k}", v] for k, v in content.by_significance.items()],
    ]
    _write_rows(summary, ["Field", "Value"], rows)
    _write_rows(
        workbook.create_sheet("Changes"), list(_CHANGE_COLUMNS), [_row(c) for c in content.changes]
    )
    _write_rows(
        workbook.create_sheet("Anomalies"),
        list(_CHANGE_COLUMNS),
        [_row(c) for c in content.anomalies],
    )
    unusual = {a["date"] for a in content.volume_anomalies}
    trend = workbook.create_sheet("Trend")
    _write_rows(
        trend,
        ["Day", "Changes", "Unusual"],
        [[day, count, "yes" if day in unusual else ""] for day, count in content.daily_trend],
    )
    if content.daily_trend:
        chart = BarChart()
        chart.title = "Changes per day"
        chart.y_axis.title = "Changes"
        chart.legend = None
        chart.add_data(
            Reference(trend, min_col=2, min_row=1, max_row=trend.max_row), titles_from_data=True
        )
        chart.set_categories(Reference(trend, min_col=1, min_row=2, max_row=trend.max_row))
        chart.width, chart.height = 24, 9
        trend.add_chart(chart, "E2")
    _write_rows(workbook.create_sheet("Unusual days"), _UNUSUAL_COLUMNS, _unusual_rows(content))
    _write_rows(
        workbook.create_sheet("Alerts"),
        ["Triggered", "Severity", "Status", "Title"],
        [[a["triggered_at"], a["severity"], a["status"], a["title"]] for a in content.alerts],
    )
    _write_rows(
        workbook.create_sheet("Insights"),
        ["Created", "Risk", "Summary", "Recommendations"],
        [
            [
                i["created_at"],
                i.get("risk_level"),
                i["summary"],
                "; ".join(i.get("recommendations", [])),
            ]
            for i in content.insights
        ],
    )
    _write_rows(
        workbook.create_sheet("Sources"),
        ["Name", "Kind", "Location"],
        [[s["name"], s["kind"], s["location"]] for s in content.sources],
    )
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _p(text: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(
        escape(clean_text(str(text), max_length=4000, multiline=True)).replace("\n", "<br/>"), style
    )


def _table(
    header: list[str], rows: list[list[Any]], style: ParagraphStyle, widths: list[float]
) -> Table:
    data = [[_p(h, style) for h in header]] + [[_p(cell, style) for cell in row] for row in rows]
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9ca3af")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def _volume_chart(daily: list[tuple[str, int]], unusual: set[str], width: float) -> Drawing | None:
    """Bar chart of changes per day; unusual days in red. None when there is nothing to plot."""
    peak = max((count for _, count in daily), default=0)
    if len(daily) < 2 or peak == 0:
        return None
    height = 48 * mm
    chart = VerticalBarChart()
    chart.x, chart.y = round(10 * mm), round(9 * mm)
    chart.width, chart.height = round(width - 12 * mm), round(height - 13 * mm)
    chart.data = [[count for _, count in daily]]
    every = max(1, math.ceil(len(daily) / 10))  # about ten date labels, whatever the period
    chart.categoryAxis.categoryNames = [
        day[5:] if index % every == 0 else "" for index, (day, _) in enumerate(daily)
    ]
    chart.categoryAxis.labels.fontSize = 6
    chart.categoryAxis.labels.angle = 30
    chart.categoryAxis.labels.boxAnchor = "ne"
    step = max(1, math.ceil(peak / 5))
    chart.valueAxis.valueMin = 0
    chart.valueAxis.valueMax = step * math.ceil(peak / step)
    chart.valueAxis.valueStep = step
    chart.valueAxis.labels.fontSize = 6
    chart.valueAxis.labelTextFormat = "%d"
    chart.bars[0].fillColor = _BAR
    chart.bars[0].strokeColor = None
    for index, (day, _) in enumerate(daily):
        if day in unusual:
            chart.bars[(0, index)].fillColor = _BAR_UNUSUAL
    drawing = Drawing(width, height)
    drawing.add(chart)
    return drawing


def render_pdf(content: ReportContent) -> bytes:
    buffer = io.BytesIO()
    sheet = getSampleStyleSheet()
    styles = {name: cast(ParagraphStyle, sheet[name]) for name in ("Title", "Heading2", "BodyText")}
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, leading=9.5)
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=truncate(content.title, 100),
        author="NexusFlow AI",
    )
    story: list[Any] = [
        _p(content.title, styles["Title"]),
        _p(
            f"{content.organization} / {content.project} / {content.dataset or 'all datasets'} - "
            f"{content.period_start:%Y-%m-%d} to {content.period_end:%Y-%m-%d} "
            f"(generated {content.generated_at:%Y-%m-%d %H:%M} UTC)",
            body,
        ),
        Spacer(1, 6 * mm),
        _p("Executive summary", styles["Heading2"]),
        _p(content.executive_summary, body),
        Spacer(1, 4 * mm),
        _p("Totals", styles["Heading2"]),
        _table(
            ["Metric", "Value"],
            [
                [k, v]
                for k, v in {
                    **content.totals,
                    **{f"significance {k}": v for k, v in content.by_significance.items()},
                }.items()
            ],
            body,
            [80 * mm, 40 * mm],
        ),
        Spacer(1, 4 * mm),
        _p("Trend", styles["Heading2"]),
        _p(content.trend_note or "No trend information.", body),
    ]
    unusual = {a["date"] for a in content.volume_anomalies}
    chart = _volume_chart(content.daily_trend, unusual, document.width)
    if chart is not None:
        story.extend([Spacer(1, 2 * mm), chart])
        if unusual:
            story.append(_p("Red bars: days with unusually many changes.", small))
    sections: list[tuple[str, list[str], list[list[Any]], list[float]]] = [
        (
            "Unusual activity",
            _UNUSUAL_COLUMNS,
            _unusual_rows(content),
            [35 * mm, 25 * mm, 40 * mm, 35 * mm],
        ),
        (
            "Anomalies: high and critical changes",
            ["Detected", "Record", "Type", "Level", "Changes"],
            [
                [
                    c["detected_at"][:16],
                    c["record_key"],
                    c["change_type"],
                    c["significance"],
                    _diff_summary(c["diff"]),
                ]
                for c in content.anomalies
            ],
            [28 * mm, 40 * mm, 18 * mm, 16 * mm, 78 * mm],
        ),
        (
            "Top changes",
            ["Detected", "Record", "Type", "Score", "Changes"],
            [
                [
                    c["detected_at"][:16],
                    c["record_key"],
                    c["change_type"],
                    c["score"],
                    _diff_summary(c["diff"]),
                ]
                for c in content.changes[:50]
            ],
            [28 * mm, 40 * mm, 18 * mm, 14 * mm, 80 * mm],
        ),
        (
            "Important events: alerts",
            ["Triggered", "Severity", "Status", "Title"],
            [
                [a["triggered_at"][:16], a["severity"], a["status"], a["title"]]
                for a in content.alerts[:50]
            ],
            [30 * mm, 20 * mm, 25 * mm, 105 * mm],
        ),
        (
            "AI insights",
            ["Created", "Risk", "Summary"],
            [[i["created_at"][:16], i.get("risk_level"), i["summary"]] for i in content.insights],
            [30 * mm, 18 * mm, 132 * mm],
        ),
        (
            "Sources",
            ["Name", "Kind", "Location"],
            [[s["name"], s["kind"], s["location"]] for s in content.sources],
            [45 * mm, 25 * mm, 110 * mm],
        ),
    ]
    for title, header, rows, widths in sections:
        if rows:
            story.extend(
                [
                    Spacer(1, 4 * mm),
                    _p(title, styles["Heading2"]),
                    _table(header, rows, small, widths),
                ]
            )
    document.build(story)
    return buffer.getvalue()


class ReportRendererRegistry:
    def render(self, content: ReportContent, fmt: ReportFormat) -> bytes:
        match fmt:
            case ReportFormat.JSON:
                return render_json(content)
            case ReportFormat.CSV:
                return render_csv(content)
            case ReportFormat.XLSX:
                return render_xlsx(content)
            case ReportFormat.PDF:
                return render_pdf(content)
