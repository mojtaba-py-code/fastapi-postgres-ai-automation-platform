"""Hostile CSV / XLSX payloads for the upload security tests, built in memory.

Archives are written record by record (``RawEntry`` -> local header, data,
central directory) instead of through ``zipfile``, so a test controls exactly
what an attacker controls: entry names byte for byte (``zipfile`` rewrites
backslashes on Windows), general-purpose flags (the encryption bit), Unix mode
bits (symlinks) and a compressed stream that disagrees with the declared sizes.
The other workbook parts come from a genuine ``openpyxl`` workbook; the
worksheet is hand-written so every test knows the exact XML it corrupts.

No real malware is used: the fake scanners "detect" the harmless ``MARKER``,
so local antivirus software never quarantines the suite.
"""

from __future__ import annotations

import io
import struct
import zipfile
import zlib
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from functools import cache

import openpyxl
from openpyxl.utils import get_column_letter

from nexusflow.domain.uploads.model import MAX_XLSX_ENTRIES, MAX_XLSX_UNCOMPRESSED

MARKER = b"NEXUSFLOW-HARMLESS-TEST-MARKER"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
SHEET = "xl/worksheets/sheet1.xml"
CONTENT_TYPES = "[Content_Types].xml"
WORKBOOK_RELS = "xl/_rels/workbook.xml.rels"
WORKBOOK_MAIN_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
MACRO_ENABLED_MAIN_TYPE = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
VBA_PROJECT_TYPE = "application/vnd.ms-office.vbaProject"
VBA_PROJECT_REL = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
# Header of an OLE2 / Compound File Binary document: legacy .xls, VBA projects
# and Office-encrypted OOXML packages ("EncryptedPackage") all start with it.
CFB_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
CFB_DOCUMENT = CFB_SIGNATURE + bytes(504)
# Members up to this uncompressed size are exempt from the compression-ratio check.
RATIO_EXEMPT_SIZE = 1_000_000

_LOCAL_HEADER = struct.Struct("<IHHHHHIIIHH")
_CENTRAL_HEADER = struct.Struct("<IHHHHHHIIIHHHHHII")
_END_OF_DIRECTORY = struct.Struct("<IHHHHIIH")
_DOS_DATE = (1 << 5) | 1  # 1980-01-01, the first date a ZIP header can hold
_FLAG_UTF8 = 0x800
_MADE_BY_UNIX = (3 << 8) | 20


@dataclass(frozen=True, slots=True)
class RawEntry:
    """One archive member exactly as written.

    ``stream`` replaces the compressed bytes while the declared sizes and CRC
    still describe ``data``: a member whose headers lie about what it inflates to.
    """

    name: str
    data: bytes
    deflate: bool = True
    flag_bits: int = 0
    unix_mode: int | None = None
    stream: bytes | None = None


def deflate_stream(*chunks: bytes) -> bytes:
    """A raw DEFLATE stream of ``chunks``, compressed one chunk at a time."""
    compressor = zlib.compressobj(6, zlib.DEFLATED, -15)
    return b"".join(compressor.compress(chunk) for chunk in chunks) + compressor.flush()


def build_zip(entries: Iterable[RawEntry]) -> bytes:
    out = io.BytesIO()
    directory: list[bytes] = []
    # Members sharing one ``data`` object (padding) are compressed only once; the
    # memo keeps each object alive, so its ``id`` cannot be reused meanwhile.
    memo: dict[tuple[int, bool], tuple[bytes, bytes, int]] = {}
    for entry in entries:
        name = entry.name.encode("utf-8")
        flags = entry.flag_bits | (0 if name.isascii() else _FLAG_UTF8)
        key = (id(entry.data), entry.deflate)
        if key not in memo:
            packed = deflate_stream(entry.data) if entry.deflate else entry.data
            memo[key] = (entry.data, packed, zlib.crc32(entry.data))
        _, packed, crc = memo[key]
        payload = entry.stream if entry.stream is not None else packed
        method = zipfile.ZIP_DEFLATED if entry.deflate else zipfile.ZIP_STORED
        made_by = 20 if entry.unix_mode is None else _MADE_BY_UNIX
        external = 0 if entry.unix_mode is None else entry.unix_mode << 16
        sizes = (crc, len(payload), len(entry.data), len(name))
        offset = out.tell()
        out.write(_LOCAL_HEADER.pack(0x04034B50, 20, flags, method, 0, _DOS_DATE, *sizes, 0))
        out.write(name + payload)
        directory.append(
            _CENTRAL_HEADER.pack(
                0x02014B50,
                made_by,
                20,
                flags,
                method,
                0,
                _DOS_DATE,
                *sizes,
                0,
                0,
                0,
                0,
                external,
                offset,
            )
            + name
        )
    start = out.tell()
    out.write(b"".join(directory))
    size = out.tell() - start
    count = len(directory)
    out.write(_END_OF_DIRECTORY.pack(0x06054B50, 0, 0, count, count, size, start, 0))
    return out.getvalue()


# ------------------------------------------------------------------ workbooks


def _xml_text(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def inline_row(number: int, values: Sequence[str]) -> str:
    """An XML ``<row>`` of inline-string cells, starting in column A."""
    cells = "".join(
        f'<c r="{get_column_letter(column)}{number}" t="inlineStr">'
        f"<is><t>{_xml_text(value)}</t></is></c>"
        for column, value in enumerate(values, start=1)
    )
    return f'<row r="{number}">{cells}</row>'


def sheet_xml(rows: str, *, prolog: str = "") -> bytes:
    return (
        f'{prolog}<worksheet xmlns="{MAIN_NS}"><sheetData>{rows}</sheetData></worksheet>'.encode()
    )


HEADER_ROW = inline_row(1, ["SKU", "Title"])
DEFAULT_SHEET = sheet_xml(HEADER_ROW + inline_row(2, ["A-1", "Widget"]))


@cache
def _openpyxl_parts() -> tuple[tuple[str, bytes], ...]:
    buffer = io.BytesIO()
    openpyxl.Workbook().save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        return tuple((info.filename, archive.read(info)) for info in archive.infolist())


def workbook_parts(sheet: bytes = DEFAULT_SHEET) -> dict[str, bytes]:
    """The parts of a genuine one-sheet workbook whose worksheet XML is ``sheet``."""
    parts = dict(_openpyxl_parts())
    parts[SHEET] = sheet
    return parts


def workbook(
    parts: Mapping[str, bytes] | None = None,
    *,
    extra: Iterable[RawEntry] = (),
    overrides: Mapping[str, RawEntry] | None = None,
) -> bytes:
    """An XLSX archive of ``parts`` (default: a valid workbook) plus ``extra`` members;
    ``overrides`` replaces the archive record written for a part."""
    chosen = workbook_parts() if parts is None else parts
    records = overrides or {}
    members = [records.get(name, RawEntry(name, data)) for name, data in chosen.items()]
    return build_zip([*members, *extra])


def add_override(parts: dict[str, bytes], part_name: str, content_type: str) -> None:
    override = f'<Override PartName="{part_name}" ContentType="{content_type}"/></Types>'
    parts[CONTENT_TYPES] = parts[CONTENT_TYPES].replace(b"</Types>", override.encode())


def add_relationship(
    parts: dict[str, bytes], rel_type: str, target: str, rel_id: str = "rIdHostile"
) -> None:
    relationship = f'<Relationship Id="{rel_id}" Type="{rel_type}" Target="{target}"/>'
    closing = b"</Relationships>"
    parts[WORKBOOK_RELS] = parts[WORKBOOK_RELS].replace(closing, relationship.encode() + closing)


def macro_enabled_parts() -> dict[str, bytes]:
    """A workbook whose main part is declared macro-enabled (what Excel saves as .xlsm)."""
    parts = workbook_parts()
    assert WORKBOOK_MAIN_TYPE.encode() in parts[CONTENT_TYPES]
    parts[CONTENT_TYPES] = parts[CONTENT_TYPES].replace(
        WORKBOOK_MAIN_TYPE.encode(), MACRO_ENABLED_MAIN_TYPE.encode()
    )
    return parts


# ------------------------------------------------------------ hostile archives


def ratio_bomb() -> bytes:
    """A valid worksheet padded with 4 MiB of whitespace: about 1000:1 compression."""
    padded = DEFAULT_SHEET.replace(b"</sheetData>", b"</sheetData>" + b" " * (4 << 20))
    return workbook(workbook_parts(padded))


def total_size_bomb() -> bytes:
    """Members of exactly ``RATIO_EXEMPT_SIZE`` zero bytes (none trips the ratio
    check on its own) that together inflate past ``MAX_XLSX_UNCOMPRESSED``."""
    count = MAX_XLSX_UNCOMPRESSED // RATIO_EXEMPT_SIZE + 1
    padding = bytes(RATIO_EXEMPT_SIZE)
    return workbook(extra=[RawEntry(f"xl/media/pad{i}.bin", padding) for i in range(count)])


def with_entry_count(total: int) -> bytes:
    base = workbook_parts()
    filler = [RawEntry(f"xl/media/m{i}.bin", b"", deflate=False) for i in range(total - len(base))]
    return workbook(base, extra=filler)


def too_many_entries() -> bytes:
    return with_entry_count(MAX_XLSX_ENTRIES + 1)


def encrypted_sheet() -> bytes:
    """The worksheet member carries the (PKWARE) encryption flag, general-purpose bit 0."""
    return workbook(overrides={SHEET: RawEntry(SHEET, DEFAULT_SHEET, flag_bits=0x1)})


def with_member(name: str, data: bytes = b"<x/>", *, unix_mode: int | None = None) -> bytes:
    return workbook(extra=[RawEntry(name, data, unix_mode=unix_mode)])


def vba_project() -> bytes:
    return with_member("xl/vbaProject.bin", CFB_DOCUMENT)


def lying_sheet(padding: int = 4 << 20) -> bytes:
    """The worksheet declares its true size and CRC, but its DEFLATE stream goes on to
    inflate ``padding`` NUL bytes more - which would break the XML if ever produced."""
    stream = deflate_stream(DEFAULT_SHEET, bytes(padding))
    return workbook(overrides={SHEET: RawEntry(SHEET, DEFAULT_SHEET, stream=stream)})


def sparse_sheet(last_row: int) -> bytes:
    """A header row and one data row numbered ``last_row``, nothing in between and
    no ``<dimension>``: a few kilobytes whatever ``last_row`` is."""
    return workbook(workbook_parts(sheet_xml(HEADER_ROW + inline_row(last_row, ["A-1", "W"]))))


def csv_bytes() -> bytes:
    return b"SKU,Title\r\nA-1,Widget\r\n"


def truncated_workbook() -> bytes:
    return workbook()[:300]


def empty_file() -> bytes:
    return b""


def cfb_document() -> bytes:
    return CFB_DOCUMENT
