"""Change analytics: totals from aggregated counts, unusual days and the trend."""

from __future__ import annotations

import io
import json
import zipfile
from collections import Counter
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta, timezone

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.graphics.charts.barcharts import VerticalBarChart

from nexusflow.domain.records.model import ChangeType, ChangeVolume, Significance
from nexusflow.domain.reports.analytics import (
    every_day,
    headline,
    summarize_volume,
    trend_note,
    volume_spikes,
)
from nexusflow.domain.reports.model import ReportContent, ReportFormat
from nexusflow.infrastructure.reporting.renderers import (
    _BAR,
    _BAR_UNUSUAL,
    ReportRendererRegistry,
    _volume_chart,
)


def _days(*counts: int) -> list[tuple[str, int]]:
    return [(f"2026-03-{day:02d}", count) for day, count in enumerate(counts, start=1)]


class TestVolumeSpikes:
    def test_a_burst_day_is_flagged_with_its_baseline(self) -> None:
        spikes = volume_spikes(_days(4, 6, 5, 7, 5, 48, 6, 5))
        assert [s["date"] for s in spikes] == ["2026-03-06"]
        assert spikes[0]["changes"] == 48
        assert spikes[0]["baseline"] == 5.5
        assert spikes[0]["score"] > 3.5

    def test_ordinary_variation_is_not_an_anomaly(self) -> None:
        assert volume_spikes(_days(10, 14, 9, 12, 15, 11, 13, 10)) == []

    def test_small_counts_are_never_flagged(self) -> None:
        # 4 changes on a quiet dataset is 4x the norm, but still noise.
        assert volume_spikes(_days(0, 1, 0, 1, 4, 0, 1)) == []

    def test_a_flat_baseline_still_reveals_a_spike(self) -> None:
        spikes = volume_spikes(_days(0, 0, 0, 0, 9, 0, 0))
        assert [(s["date"], s["score"]) for s in spikes] == [("2026-03-05", None)]

    def test_one_extreme_day_does_not_hide_another(self) -> None:
        spikes = volume_spikes(_days(5, 4, 6, 400, 5, 60, 4, 6, 5, 5))
        assert {s["date"] for s in spikes} == {"2026-03-04", "2026-03-06"}

    def test_short_periods_have_no_baseline(self) -> None:
        assert volume_spikes(_days(1, 90, 2)) == []


class TestTrendNote:
    @pytest.mark.parametrize(
        ("counts", "expected"),
        [
            ((10, 10, 20, 20), "Change volume rose 100% from the first to the second half."),
            ((20, 20, 10, 10), "Change volume fell 50% from the first to the second half."),
            ((10, 10, 10, 10), "Change volume was stable across the period."),
            ((0, 0, 0, 0), "No changes were detected in the period."),
            ((0, 0, 3, 4), "All detected changes fall in the second half of the period."),
            ((7,), "Not enough days in the period to describe a trend."),
            # An odd number of days: the halves differ in length, so their daily
            # averages are compared - a flat week is flat, not "rose 33%".
            ((10,) * 7, "Change volume was stable across the period."),
            (
                (10, 10, 10, 20, 20, 20, 20),
                "Change volume rose 100% from the first to the second half.",
            ),
            ((9, 9, 3, 3, 3), "Change volume fell 67% from the first to the second half."),
        ],
    )
    def test_the_sentence_matches_the_numbers(self, counts: tuple[int, ...], expected: str) -> None:
        assert trend_note(_days(*counts)) == expected


class TestEveryDay:
    def test_quiet_days_are_included_and_the_end_is_exclusive(self) -> None:
        start = datetime(2026, 3, 1, tzinfo=UTC)
        end = datetime(2026, 3, 4, tzinfo=UTC)  # exclusive, like the queries
        days = every_day(start, end, Counter({"2026-03-02": 3}))
        assert days == [("2026-03-01", 0), ("2026-03-02", 3), ("2026-03-03", 0)]

    def test_days_are_utc_whatever_the_request_offset(self) -> None:
        tehran = timezone(timedelta(hours=3, minutes=30))
        start = datetime(2026, 3, 1, 2, 0, tzinfo=tehran)  # 2026-02-28 22:30 UTC
        end = datetime(2026, 3, 2, 2, 0, tzinfo=tehran)
        days = every_day(start, end, Counter())
        assert [d for d, _ in days] == ["2026-02-28", "2026-03-01"]

    def test_the_series_is_bounded(self) -> None:
        start = datetime(2020, 1, 1, tzinfo=UTC)
        end = datetime(2026, 1, 1, tzinfo=UTC)
        assert len(every_day(start, end, Counter())) == 367

    def test_a_longest_period_off_midnight_keeps_its_last_day(self) -> None:
        start = datetime(2026, 1, 1, 12, tzinfo=UTC)
        end = start + timedelta(days=366)  # the longest period allowed, spanning 367 days
        days = every_day(start, end, Counter({"2027-01-02": 4}))
        assert len(days) == 367
        assert days[-1] == ("2027-01-02", 4)


def _volume(day: int, kind: ChangeType, level: Significance, count: int) -> ChangeVolume:
    return ChangeVolume(day=date(2026, 3, day), change_type=kind, significance=level, count=count)


class TestSummarizeVolume:
    START = datetime(2026, 3, 1, tzinfo=UTC)
    END = datetime(2026, 3, 9, tzinfo=UTC)

    def test_totals_add_up_every_group(self) -> None:
        summary = summarize_volume(
            [
                _volume(2, ChangeType.UPDATED, Significance.LOW, 4000),
                _volume(2, ChangeType.UPDATED, Significance.HIGH, 1500),
                _volume(5, ChangeType.CREATED, Significance.MEDIUM, 700),
                _volume(5, ChangeType.DELETED, Significance.MEDIUM, 3),
            ],
            self.START,
            self.END,
        )
        # Beyond any listing cap: counts come from the database, not from a sample.
        assert summary.totals == {"changes": 6203, "created": 700, "updated": 5500, "deleted": 3}
        assert summary.by_significance == {"low": 4000, "medium": 703, "high": 1500, "critical": 0}
        assert dict(summary.daily)["2026-03-02"] == 5500
        assert dict(summary.daily)["2026-03-05"] == 703
        assert len(summary.daily) == 8  # quiet days included

    def test_an_empty_period_has_a_stable_shape(self) -> None:
        summary = summarize_volume([], self.START, self.END)
        assert summary.totals == {"changes": 0, "created": 0, "updated": 0, "deleted": 0}
        assert summary.by_significance == {"low": 0, "medium": 0, "high": 0, "critical": 0}
        assert summary.unusual_days == []
        assert summary.trend_note == "No changes were detected in the period."

    def test_the_trend_and_unusual_days_follow_the_series(self) -> None:
        counts = (4, 6, 5, 7, 5, 48, 6, 5)
        rows = [
            _volume(day, ChangeType.UPDATED, Significance.LOW, count)
            for day, count in enumerate(counts, start=1)
        ]
        summary = summarize_volume(rows, self.START, self.END)
        assert [day["date"] for day in summary.unusual_days] == ["2026-03-06"]
        assert summary.trend_note == trend_note(_days(*counts))


class TestHeadline:
    START = datetime(2026, 3, 1, tzinfo=UTC)
    END = datetime(2026, 3, 9, tzinfo=UTC)

    def test_it_counts_every_change_and_names_the_highest_level(self) -> None:
        summary = summarize_volume(
            [
                _volume(2, ChangeType.UPDATED, Significance.LOW, 4000),
                _volume(3, ChangeType.UPDATED, Significance.HIGH, 2),
                _volume(3, ChangeType.CREATED, Significance.MEDIUM, 9),
            ],
            self.START,
            self.END,
        )
        assert headline(summary, "Prices") == (
            "4011 changes detected in 'Prices': 4002 updated, 9 created, 0 removed. "
            "Highest significance: high."
        )

    def test_one_change_is_singular(self) -> None:
        summary = summarize_volume(
            [_volume(2, ChangeType.DELETED, Significance.MEDIUM, 1)], self.START, self.END
        )
        assert headline(summary, "Prices").startswith("1 change detected in 'Prices'")

    def test_a_quiet_period_says_so(self) -> None:
        summary = summarize_volume([], self.START, self.END)
        assert headline(summary, "Prices") == "No changes were detected in 'Prices' in this period."


def _content() -> ReportContent:
    daily = _days(4, 6, 5, 7, 5, 48, 6, 5)
    return ReportContent(
        title="Monthly report",
        organization="Org",
        project="Proj",
        dataset=None,
        period_start=datetime(2026, 3, 1, tzinfo=UTC),
        period_end=datetime(2026, 3, 9, tzinfo=UTC),
        generated_at=datetime(2026, 3, 9, tzinfo=UTC),
        executive_summary="summary",
        totals={"changes": 86},
        by_significance={"low": 86},
        daily_trend=daily,
        changes=[],
        anomalies=[],
        alerts=[],
        insights=[],
        sources=[],
        volume_anomalies=volume_spikes(daily),
        trend_note=trend_note(daily),
    )


class TestRenderedTrend:
    def test_json_carries_the_trend_and_the_unusual_days(self) -> None:
        data = json.loads(ReportRendererRegistry().render(_content(), ReportFormat.JSON))
        assert data["trend_note"].startswith("Change volume rose")
        assert [a["date"] for a in data["volume_anomalies"]] == ["2026-03-06"]

    def test_xlsx_flags_unusual_days_and_charts_the_trend(self) -> None:
        data = ReportRendererRegistry().render(_content(), ReportFormat.XLSX)
        workbook = openpyxl.load_workbook(io.BytesIO(data))
        trend = workbook["Trend"]
        assert isinstance(trend, Worksheet)
        flagged = [row[0] for row in trend.iter_rows(min_row=2, values_only=True) if row[2]]
        assert flagged == ["2026-03-06"]
        unusual = list(workbook["Unusual days"].iter_rows(min_row=2, values_only=True))
        assert unusual == [("2026-03-06", 48, 5.5, unusual[0][3])]
        with zipfile.ZipFile(io.BytesIO(data)) as archive:  # a native Excel chart
            assert "xl/charts/chart1.xml" in archive.namelist()

    @pytest.mark.parametrize(
        "daily",
        [
            _days(4, 6, 5, 7, 5, 48, 6, 5),
            _days(0, 0, 0),  # nothing to plot
            _days(3),  # a single day
            [(f"2025-{m:02d}-{d:02d}", (m * d) % 17) for m in range(1, 13) for d in range(1, 29)],
        ],
    )
    def test_pdf_renders_any_trend_shape(self, daily: list[tuple[str, int]]) -> None:
        content = replace(
            _content(),
            daily_trend=daily,
            volume_anomalies=volume_spikes(daily),
            trend_note=trend_note(daily),
        )
        assert ReportRendererRegistry().render(content, ReportFormat.PDF).startswith(b"%PDF")

    def test_the_pdf_chart_marks_unusual_days_and_thins_its_labels(self) -> None:
        daily = _days(*([5] * 25), 60, 5, 5, 5, 5)
        drawing = _volume_chart(daily, {"2026-03-26"}, 180.0)
        assert drawing is not None
        chart = drawing.contents[0]
        assert isinstance(chart, VerticalBarChart)
        assert chart.bars[(0, 25)].fillColor == _BAR_UNUSUAL
        assert chart.bars[(0, 24)].fillColor == _BAR
        labels = [name for name in chart.categoryAxis.categoryNames if name]
        assert len(labels) <= 10
        assert labels[0] == "03-01"
